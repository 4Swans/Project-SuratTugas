import os
from datetime import date

import docx

from app.services.dokumen_service import generate_surat_tugas_docx

TEMPLATE_PATH = "app/templates/surat_tugas.docx"


def test_generate_1_pegawai(tmp_path):
    output_path = str(tmp_path / "surat_1.docx")
    p1 = {
        "nama": "Oki Suryowahono, S.T., M.T.",
        "nip": "197909072008031001",
        "pangkat_gol": "Pembina Tingkat I / IVc",
        "jabatan": "Kepala Pusat Kebijakan Strategis",
    }
    generate_surat_tugas_docx(
        template_path=TEMPLATE_PATH,
        output_path=output_path,
        nomor_surat="864/SJ.9/KP.01.06/09/2026",
        dasar_surat="Undangan Sekretaris Direktorat Jenderal Nomor B-2845",
        perihal="Rapat Pembahasan Kebijakan",
        tanggal_mulai=date(2026, 9, 23),
        tanggal_selesai=date(2026, 9, 25),
        lokasi="Bandung, Jawa Barat",
        penandatangan_nama="Oki Suryowahono, S.T., M.T.",
        penandatangan_jabatan="Kepala Pusat Kebijakan Strategis",
        pegawai_list=[p1],
    )

    assert os.path.exists(output_path)
    doc = docx.Document(output_path)
    # Harus ada 1 tabel utama saja (tidak ada tabel lampiran)
    assert len(doc.tables) == 1
    t0 = doc.tables[0]
    # Cek baris Kepada terisi nama
    assert "Oki Suryowahono" in t0.rows[3].cells[3].text


def test_generate_multi_pegawai_lampiran(tmp_path):
    output_path = str(tmp_path / "surat_lampiran.docx")
    pegawai_list = [
        {
            "nama": f"Pegawai {i}",
            "nip": f"19900101202001100{i}",
            "pangkat_gol": "Penata / IIIc",
            "jabatan": f"Jabatan {i}",
        }
        for i in range(1, 6)  # 5 pegawai
    ]
    generate_surat_tugas_docx(
        template_path=TEMPLATE_PATH,
        output_path=output_path,
        nomor_surat="865/SJ.9/KP.01.06/09/2026",
        dasar_surat="Undangan Nomor 456",
        perihal="Rapat Koordinasi Bersama",
        tanggal_mulai=date(2026, 9, 23),
        tanggal_selesai=date(2026, 9, 23),
        lokasi="Jakarta",
        penandatangan_nama="Oki Suryowahono, S.T., M.T.",
        penandatangan_jabatan="Kepala Pusat Kebijakan Strategis",
        pegawai_list=pegawai_list,
    )

    assert os.path.exists(output_path)
    doc = docx.Document(output_path)
    # Harus ada 2 tabel (tabel utama + tabel lampiran)
    assert len(doc.tables) == 2

    # Tabel 0 harus tertulis '(Daftar Terlampir)'
    t0 = doc.tables[0]
    assert "(Daftar Terlampir)" in t0.rows[3].cells[2].text

    # Tabel lampiran harus memiliki 1 baris header + 5 baris pegawai = 6 baris
    t_lamp = doc.tables[1]
    assert len(t_lamp.rows) == 6
    assert t_lamp.rows[0].cells[1].text == "Nama"
    assert t_lamp.rows[1].cells[1].text == "Pegawai 1"
    assert t_lamp.rows[5].cells[1].text == "Pegawai 5"
