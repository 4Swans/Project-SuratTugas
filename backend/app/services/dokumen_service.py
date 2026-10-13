"""
Service untuk pengisian template dokumen Word (.docx) surat tugas.
Preservasi 100% format naskah dinas resmi:
- Mempertahankan tab stops, indents, spacing, alignments, dan run properties dari template.
- Kasus 1 pegawai: mengisi baris Nama, NIP, Pangkat/Gol, Jabatan secara in-place.
- Kasus 2-3 pegawai: menggandakan baris 3-6 dengan deepcopy XML persis format template.
- Kasus >3 pegawai: baris Kepada diisi '(Daftar Terlampir)', baris kosong dihapus,
  serta membuat Lampiran I lengkap dengan tabel 5 kolom berstandar resmi.
"""

import copy
import os
from datetime import date
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt

NAMA_BULAN = [
    "",
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]


def format_tanggal_indonesia(d: date) -> str:
    return f"{d.day} {NAMA_BULAN[d.month]} {d.year}"


def format_periode_tanggal(tgl_mulai: date, tgl_selesai: date) -> str:
    if tgl_mulai == tgl_selesai:
        return format_tanggal_indonesia(tgl_mulai)
    if tgl_mulai.year == tgl_selesai.year:
        if tgl_mulai.month == tgl_selesai.month:
            return (
                f"{tgl_mulai.day} - {tgl_selesai.day} "
                f"{NAMA_BULAN[tgl_mulai.month]} {tgl_mulai.year}"
            )
        return (
            f"{tgl_mulai.day} {NAMA_BULAN[tgl_mulai.month]} - "
            f"{tgl_selesai.day} {NAMA_BULAN[tgl_selesai.month]} {tgl_mulai.year}"
        )
    return f"{format_tanggal_indonesia(tgl_mulai)} - {format_tanggal_indonesia(tgl_selesai)}"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f"<w:tblBorders {nsdecls('w')}>\n"
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def apply_font_to_cell(cell, sz_val: int = 20):
    """Set font size in half-points (20 = 10pt) across all paragraphs and runs in a cell."""
    for p in cell.paragraphs:
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            sz = rPr.find(qn("w:sz"))
            if sz is None:
                sz = OxmlElement("w:sz")
                rPr.append(sz)
            sz.set(qn("w:val"), str(sz_val))
            szCs = rPr.find(qn("w:szCs"))
            if szCs is None:
                szCs = OxmlElement("w:szCs")
                rPr.append(szCs)
            szCs.set(qn("w:val"), str(sz_val))


def format_signature_paragraph(p, text: str, left_dxa: int = 5200):
    """Format right-aligned signature paragraph cleanly without stray spaces or awkward center alignment."""
    update_single_run_paragraph(p, text)
    pPr = p._p.get_or_add_pPr()
    jc = pPr.find(qn("w:jc"))
    if jc is not None:
        pPr.remove(jc)
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:left"), str(left_dxa))
    ind.set(qn("w:right"), "0")
    if qn("w:firstLine") in ind.attrib:
        del ind.attrib[qn("w:firstLine")]


def convert_docx_to_pdf(docx_path: str, pdf_path: str) -> str:
    """
    Mengonversi file .docx ke .pdf menggunakan Microsoft Word COM Automation.
    Akurat 100% mempertahankan tata letak, margin, font Arial, dan header surat.
    """
    import faulthandler

    import pythoncom
    import win32com.client

    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"File docx sumber tidak ditemukan: {docx_path}")

    os.makedirs(os.path.dirname(os.path.abspath(pdf_path)), exist_ok=True)
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)

    was_fh_enabled = faulthandler.is_enabled()
    if was_fh_enabled:
        faulthandler.disable()

    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(abs_docx)
        doc.SaveAs(abs_pdf, FileFormat=17)
        doc.Close(False)
        doc = None
        word.Quit()
        word = None
        return abs_pdf
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
        if was_fh_enabled:
            faulthandler.enable()


def update_single_run_paragraph(p, new_text: str):
    """Update teks paragraf pada run pertama, kosongkan run berikutnya tanpa merusak pPr/rPr."""
    if not p.runs:
        r = p.add_run(new_text)
        r.font.name = "Arial"
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def generate_surat_tugas_docx(
    template_path: str,
    output_path: str,
    nomor_surat: str,
    dasar_surat: str,
    perihal: str,
    tanggal_mulai: date,
    tanggal_selesai: date,
    lokasi: str,
    penandatangan_nama: str,
    penandatangan_jabatan: str,
    tanggal_surat: date | None = None,
    menimbang: str | None = None,
    pegawai_list: list[dict[str, Any]] | None = None,
) -> str:
    """
    Mengisi file template Word dan menyimpannya ke output_path dengan mempertahankan
    100% format naskah dinas resmi dari template-surat-tugas-blank.docx (tab, font, spasi, indents).
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template surat tidak ditemukan: {template_path}")

    pegawai_list = pegawai_list or []
    pegawai_count = len(pegawai_list)

    if tanggal_surat is None:
        tanggal_surat = date.today()

    tgl_surat_str = format_tanggal_indonesia(tanggal_surat)
    periode_str = format_periode_tanggal(tanggal_mulai, tanggal_selesai)
    tahun_surat = tanggal_mulai.year

    if not menimbang:
        menimbang = (
            "bahwa dalam rangka mendukung pelaksanaan tugas dan fungsi Pusat Kebijakan "
            "Strategis serta menindaklanjuti undangan yang diterima, perlu menugaskan "
            "pegawai untuk mengikuti kegiatan dimaksud."
        )

    doc = docx.Document(template_path)

    if pegawai_count in (2, 3):
        sec = doc.sections[0]
        sec.top_margin = Inches(0.65)
        sec.bottom_margin = Inches(0.55)

    # 1. Update Nomor Surat di P2
    p_no = doc.paragraphs[2]
    if pegawai_count in (2, 3):
        p_surat = doc.paragraphs[1]
        p_surat.paragraph_format.space_before = Pt(0)
        p_surat.paragraph_format.space_after = Pt(4)
        p_no.paragraph_format.space_before = Pt(0)
        p_no.paragraph_format.space_after = Pt(6)

    if len(p_no.runs) >= 3:
        p_no.runs[0].text = "Nomor"
        p_no.runs[1].text = ": "
        p_no.runs[2].text = nomor_surat
        for r in p_no.runs[3:]:
            r.text = ""
    else:
        p_no.text = f"Nomor: {nomor_surat}"
        p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_no.runs:
            p_no.runs[0].font.name = "Arial"

    # 2. Update Table 0 (Menimbang, Dasar, Memberi Tugas, Kepada, Untuk)
    if doc.tables:
        t = doc.tables[0]

        # Row 0: Menimbang
        p_menimbang = t.rows[0].cells[2].paragraphs[0]
        update_single_run_paragraph(p_menimbang, menimbang)

        # Row 1: Dasar
        p_dasar = t.rows[1].cells[2].paragraphs[0]
        update_single_run_paragraph(p_dasar, dasar_surat)

        # Row 3 to 6: Kepada
        if pegawai_count == 1:
            p1 = pegawai_list[0]
            # Nama
            p_nama = t.rows[3].cells[3].paragraphs[0]
            if len(p_nama.runs) >= 2:
                p_nama.runs[0].text = ": "
                p_nama.runs[1].text = p1["nama"]
                for r in p_nama.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nama, f": {p1['nama']}")

            # NIP
            p_nip = t.rows[4].cells[3].paragraphs[0]
            if len(p_nip.runs) >= 2:
                p_nip.runs[0].text = ": "
                p_nip.runs[1].text = str(p1["nip"])
                for r in p_nip.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nip, f": {p1['nip']}")

            # Pangkat, Gol/Ruang
            p_pangkat = t.rows[5].cells[3].paragraphs[0]
            p1_pangkat = (
                p1.get("pangkat_gol")
                or p1.get("pangkat")
                or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            )
            if len(p_pangkat.runs) >= 2:
                p_pangkat.runs[0].text = ": "
                p_pangkat.runs[1].text = str(p1_pangkat)
                for r in p_pangkat.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_pangkat, f": {p1_pangkat}")

            # Jabatan
            p_jabatan = t.rows[6].cells[3].paragraphs[0]
            if len(p_jabatan.runs) >= 2:
                p_jabatan.runs[0].text = ": "
                p_jabatan.runs[1].text = str(p1["jabatan"])
                for r in p_jabatan.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_jabatan, f": {p1['jabatan']}")

        elif pegawai_count in (2, 3):
            # Format 2-3 Pegawai: TIDAK MENGGUNAKAN NOMOR pada Nama
            p1 = pegawai_list[0]
            update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], "Nama")
            p_nama = t.rows[3].cells[3].paragraphs[0]
            if len(p_nama.runs) >= 2:
                p_nama.runs[0].text = ": "
                p_nama.runs[1].text = p1["nama"]
                for r in p_nama.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nama, f": {p1['nama']}")

            update_single_run_paragraph(t.rows[4].cells[2].paragraphs[0], "NIP")
            p_nip = t.rows[4].cells[3].paragraphs[0]
            if len(p_nip.runs) >= 2:
                p_nip.runs[0].text = ": "
                p_nip.runs[1].text = str(p1["nip"])
                for r in p_nip.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nip, f": {p1['nip']}")

            update_single_run_paragraph(t.rows[5].cells[2].paragraphs[0], "Pangkat, Gol/Ruang")
            p_pangkat = t.rows[5].cells[3].paragraphs[0]
            p1_pangkat = (
                p1.get("pangkat_gol")
                or p1.get("pangkat")
                or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            )
            if len(p_pangkat.runs) >= 2:
                p_pangkat.runs[0].text = ": "
                p_pangkat.runs[1].text = str(p1_pangkat)
                for r in p_pangkat.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_pangkat, f": {p1_pangkat}")

            update_single_run_paragraph(t.rows[6].cells[2].paragraphs[0], "Jabatan")
            p_jabatan = t.rows[6].cells[3].paragraphs[0]
            if len(p_jabatan.runs) >= 2:
                p_jabatan.runs[0].text = ": "
                p_jabatan.runs[1].text = str(p1["jabatan"])
                for r in p_jabatan.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_jabatan, f": {p1['jabatan']}")

            # Kloning baris 3-6 untuk pegawai 2 (dan 3 jika ada) TANPA NOMOR
            curr_anchor = t.rows[6]
            for p_idx, p_data in enumerate(pegawai_list[1:], start=2):
                for field_idx in range(4):
                    cloned_tr = copy.deepcopy(t.rows[3 + field_idx]._tr)
                    curr_anchor._tr.addnext(cloned_tr)
                    curr_anchor = docx.table._Row(cloned_tr, t)
                    c_cells = curr_anchor.cells

                    update_single_run_paragraph(c_cells[0].paragraphs[0], "")
                    update_single_run_paragraph(c_cells[1].paragraphs[0], "")

                    if field_idx == 0:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "Nama")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = p_data["nama"]
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(
                                c_cells[3].paragraphs[0], f": {p_data['nama']}"
                            )
                    elif field_idx == 1:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "NIP")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_data["nip"])
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(
                                c_cells[3].paragraphs[0], f": {p_data['nip']}"
                            )
                    elif field_idx == 2:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "Pangkat, Gol/Ruang")
                        p_pkt = (
                            p_data.get("pangkat_gol")
                            or p_data.get("pangkat")
                            or "Penyedia Jasa Lainnya Perorangan (PJLP)"
                        )
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_pkt)
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(c_cells[3].paragraphs[0], f": {p_pkt}")
                    elif field_idx == 3:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "Jabatan")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_data["jabatan"])
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(
                                c_cells[3].paragraphs[0], f": {p_data['jabatan']}"
                            )

            # Optimasi spasi & font untuk 2-3 pegawai agar muat 1 halaman
            p_mt = t.rows[2].cells[0].paragraphs[0]
            if p_mt._p.pPr is not None:
                sp = p_mt._p.pPr.find(qn("w:spacing"))
                if sp is not None:
                    sp.set(qn("w:before"), "40")
                    sp.set(qn("w:after"), "40")

            p_unt_sp = t.rows[-1].cells[2].paragraphs[0]
            if p_unt_sp._p.pPr is not None:
                sp = p_unt_sp._p.pPr.find(qn("w:spacing"))
                if sp is not None:
                    sp.set(qn("w:before"), "80" if pegawai_count == 3 else "100")

            for r_idx in range(3, len(t.rows) - 1):
                for cell in t.rows[r_idx].cells:
                    apply_font_to_cell(cell, sz_val=20)
                    set_cell_margins(cell, top=20, bottom=20, left=60, right=60)
                    for p in cell.paragraphs:
                        pPr = p._p.get_or_add_pPr()
                        sp = pPr.find(qn("w:spacing"))
                        if sp is None:
                            sp = OxmlElement("w:spacing")
                            pPr.append(sp)
                        sp.set(qn("w:line"), "230")
                        sp.set(qn("w:lineRule"), "auto")
                        sp.set(qn("w:before"), "0")
                        sp.set(qn("w:after"), "0")

        elif pegawai_count > 3:
            # Kasus > 3 pegawai: tulis '(Daftar Terlampir)'
            update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], "(Daftar Terlampir)")
            update_single_run_paragraph(t.rows[3].cells[3].paragraphs[0], "")
            t._tbl.remove(t.rows[6]._tr)
            t._tbl.remove(t.rows[5]._tr)
            t._tbl.remove(t.rows[4]._tr)

        # Row Untuk (baris terakhir tabel 0)
        p_untuk = t.rows[-1].cells[2].paragraphs[0]
        perihal_clean = perihal.strip().rstrip(".")
        lokasi_clean = lokasi.strip().rstrip(".")
        if perihal_clean.lower().startswith("mengikuti") or perihal_clean.lower().startswith(
            "menghadiri"
        ):
            untuk_text = f"{perihal_clean} yang dilaksanakan pada {periode_str} di {lokasi_clean}."
        else:
            untuk_text = f"Mengikuti kegiatan {perihal_clean} yang dilaksanakan pada {periode_str} di {lokasi_clean}."
        update_single_run_paragraph(p_untuk, untuk_text)

    # 3. Update Paragraf Biaya, Penutup, Tanggal, dan Tanda Tangan
    # P3: Biaya
    p_biaya = doc.paragraphs[3]
    if "Tahun Anggaran" in p_biaya.text:
        p_biaya.runs[0].text = p_biaya.runs[0].text.replace("2026", str(tahun_surat))

    # P4: Demikian...
    p_demikian = doc.paragraphs[4]
    if pegawai_count in (2, 3):
        p_biaya.paragraph_format.space_before = Pt(3)
        p_biaya.paragraph_format.space_after = Pt(0)
        p_demikian.paragraph_format.space_before = Pt(3)
        p_demikian.paragraph_format.space_after = Pt(0)
    elif p_demikian._p.pPr is not None:
        sp = p_demikian._p.pPr.find(qn("w:spacing"))
        if sp is not None:
            sp.set(qn("w:before"), "40")

    # Tangkap referensi paragraf tanda tangan SEBELUM menghapus baris kosong
    p_tgl = doc.paragraphs[6]
    p_jabatan_sign = doc.paragraphs[7]
    p_nama_sign = doc.paragraphs[14]
    clean_jabatan = penandatangan_jabatan.strip().rstrip(",")

    # Format rapi pada sisi kanan dokumen (left="5400" dxa)
    format_signature_paragraph(p_tgl, f"Jakarta, {tgl_surat_str}", left_dxa=5400)
    format_signature_paragraph(p_jabatan_sign, f"{clean_jabatan},", left_dxa=5400)
    format_signature_paragraph(p_nama_sign, penandatangan_nama, left_dxa=5400)

    # Sesuaikan jumlah baris kosong tanda tangan agar luas untuk ttd digital & stamp TU, serta tetap 1 halaman
    if pegawai_count == 1:
        # Hapus 2 baris kosong (menyisakan 4 baris kosong)
        for idx in [12, 11]:
            p = doc.paragraphs[idx]
            p._element.getparent().remove(p._element)
    elif pegawai_count in (2, 3):
        # Hapus 2 baris kosong (menyisakan 4 baris kosong yang luas untuk ttd digital & stamp TU)
        for idx in [13, 12]:
            p = doc.paragraphs[idx]
            p._element.getparent().remove(p._element)
        p5 = doc.paragraphs[5]
        if p5.text.strip() == "":
            p5._element.getparent().remove(p5._element)

    # 4. Halaman Lampiran (hanya jika > 3 pegawai)
    if pegawai_count > 3:
        doc.add_page_break()

        # Header lampiran (posisi pojok kanan atas)
        p_l1 = doc.add_paragraph()
        format_signature_paragraph(p_l1, "Lampiran I Surat Tugas", left_dxa=4600)
        p_l1.paragraph_format.space_before = Pt(0)
        p_l1.paragraph_format.space_after = Pt(2)
        p_l1.paragraph_format.line_spacing = 1.15
        p_l1.runs[0].bold = True
        p_l1.runs[0].font.name = "Arial"
        p_l1.runs[0].font.size = Pt(10)

        p_l2 = doc.add_paragraph()
        format_signature_paragraph(p_l2, f"Nomor   :  {nomor_surat}", left_dxa=4600)
        p_l2.paragraph_format.space_before = Pt(0)
        p_l2.paragraph_format.space_after = Pt(2)
        p_l2.paragraph_format.line_spacing = 1.15
        p_l2.runs[0].font.name = "Arial"
        p_l2.runs[0].font.size = Pt(10)

        p_l3 = doc.add_paragraph()
        format_signature_paragraph(p_l3, f"Tanggal : {tgl_surat_str}", left_dxa=4600)
        p_l3.paragraph_format.space_before = Pt(0)
        p_l3.paragraph_format.space_after = Pt(10)
        p_l3.paragraph_format.line_spacing = 1.15
        p_l3.runs[0].font.name = "Arial"
        p_l3.runs[0].font.size = Pt(10)

        # Tabel 5 kolom berstandar resmi
        col_widths = [Inches(0.45), Inches(1.85), Inches(1.50), Inches(1.15), Inches(1.35)]
        table_l = doc.add_table(rows=1, cols=5)
        table_l.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table_l)

        headers = ["No", "Nama", "NIP", "Pangkat/Gol\nRuang", "Jabatan"]
        hdr_cells = table_l.rows[0].cells
        for c_idx, title in enumerate(headers):
            hdr_cells[c_idx].width = col_widths[c_idx]
            set_cell_margins(hdr_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = hdr_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(title)
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.bold = True

        for idx, p_data in enumerate(pegawai_list, start=1):
            row_cells = table_l.add_row().cells
            pkt = (
                p_data.get("pangkat_gol")
                or p_data.get("pangkat")
                or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            )
            row_vals = [
                str(idx),
                p_data["nama"],
                str(p_data["nip"]),
                pkt,
                p_data["jabatan"],
            ]
            for c_idx, val in enumerate(row_vals):
                row_cells[c_idx].width = col_widths[c_idx]
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(10)

    # Pastikan direktori output ada
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
