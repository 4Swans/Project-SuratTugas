"""
Service untuk mengekstrak informasi acara dari surat undangan (PDF, Scan/Gambar, Word)
menggunakan Gemini API (SDK google-genai) dengan structured JSON output.
"""

import io
import json
import os
import time
from typing import Optional

import docx
from fastapi import HTTPException, status
from google import genai
from google.genai import types

from app.core.config import settings
from app.core.logging import logger
from app.schemas.ekstraksi import EkstraksiUndanganResponse


def get_prompt_text() -> str:
    prompt_path = os.path.join(os.path.dirname(__file__), "..", "prompts", "ekstraksi_undangan.md")
    prompt_path = os.path.abspath(prompt_path)
    if os.path.exists(prompt_path):
        with open(prompt_path, "r", encoding="utf-8") as f:
            return f.read()
    return (
        "Ekstrak informasi penting dari surat undangan: nama_acara, tanggal_mulai, "
        "tanggal_selesai, lokasi, topik, nomor_surat_undangan, instansi_pengundang, dasar_surat."
    )


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    lines = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)
    for table in doc.tables:
        for row in table.rows:
            row_vals = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_vals:
                lines.append(" | ".join(row_vals))
    return "\n".join(lines)


def extract_undangan(
    file_bytes: bytes, filename: str, content_type: Optional[str] = None
) -> EkstraksiUndanganResponse:
    if not settings.GEMINI_API_KEY:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "GEMINI_API_KEY belum dikonfigurasi pada file backend/.env. "
                "Silakan masukkan API key Gemini untuk menggunakan fitur ekstraksi undangan."
            ),
        )

    # 1. Simpan file undangan ke storage/undangan
    upload_dir = os.path.join(settings.STORAGE_DIR, "undangan")
    os.makedirs(upload_dir, exist_ok=True)
    clean_filename = f"{int(time.time())}_{os.path.basename(filename)}"
    saved_path = os.path.join(upload_dir, clean_filename)

    with open(saved_path, "wb") as f:
        f.write(file_bytes)

    # 2. Tentukan MIME type dan contents
    ext = os.path.splitext(filename)[1].lower()
    prompt_instruction = get_prompt_text()
    contents = []

    client = genai.Client(api_key=settings.GEMINI_API_KEY)

    try:
        if ext in [".docx", ".doc"]:
            # File Word: ekstraksi teks dulu
            doc_text = extract_text_from_docx(file_bytes)
            user_content = (
                f"{prompt_instruction}\n\n"
                f"--- TEKS DARI DOKUMEN UNDANGAN ({filename}) ---\n"
                f"{doc_text}"
            )
            contents = [user_content]

        elif ext == ".pdf":
            mime = "application/pdf"
            pdf_part = types.Part.from_bytes(data=file_bytes, mime_type=mime)
            contents = [prompt_instruction, pdf_part]

        elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
            mime_map = {
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".webp": "image/webp",
            }
            mime = mime_map.get(ext, "image/jpeg")
            img_part = types.Part.from_bytes(data=file_bytes, mime_type=mime)
            contents = [prompt_instruction, img_part]

        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    f"Format file '{ext}' tidak didukung. "
                    "Harap upload PDF, DOCX, atau Gambar (PNG/JPG)."
                ),
            )

        logger.info(f"Mengirim dokumen {clean_filename} ke Gemini model: {settings.GEMINI_MODEL}")

        config = types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=EkstraksiUndanganResponse,
            temperature=0.1,
        )

        response = client.models.generate_content(
            model=settings.GEMINI_MODEL,
            contents=contents,
            config=config,
        )

        response_text = response.text
        if not response_text:
            raise ValueError("Respons Gemini kosong.")

        parsed_data = json.loads(response_text)
        result = EkstraksiUndanganResponse.model_validate(parsed_data)
        result.file_name = filename
        result.file_path = saved_path

        return result

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Gagal melakukan ekstraksi dokumen {filename}: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Gagal mengekstrak data undangan dengan Gemini AI: {str(e)}",
        )
