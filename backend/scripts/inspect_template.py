import docx
import os

doc_path = 'app/templates/surat_tugas.docx'
doc = docx.Document(doc_path)
print('=== PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs):
    print(f'P{i}: text="{p.text}"')

print('=== TABLES ===')
print('Total tables:', len(doc.tables))
for t_idx, t in enumerate(doc.tables):
    print(f'Table {t_idx} rows: {len(t.rows)}, cols: {len(t.columns)}')
    for r_idx, row in enumerate(t.rows):
        cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f'  Row {r_idx}: {cells_text}')

print('=== SECTIONS & HEADER ===')
for s in doc.sections:
    print('Header paragraphs:', len(s.header.paragraphs))
    for hp in s.header.paragraphs:
        print('  Header:', hp.text)
