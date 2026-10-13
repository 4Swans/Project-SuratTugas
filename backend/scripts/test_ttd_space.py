import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import copy
from datetime import date
import docx
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pypdf
from app.services import dokumen_service

tpl_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'app', 'templates', 'surat_tugas.docx')

pegs = [
    {'nama': 'Dr. Ismail, S.Kom., M.Si.', 'nip': '197503122000031001', 'pangkat_gol': 'Pembina Utama Madya (IV/d)', 'jabatan': 'Analis Kebijakan Ahli Utama'},
    {'nama': 'Achmad Yansyuru, S.Kom.', 'nip': '198804152014021003', 'pangkat_gol': 'Penata Muda Tk. I (III/b)', 'jabatan': 'Pranata Komputer Ahli Muda'},
    {'nama': 'Priska Apnitami, S.Stat.', 'nip': '199510102020122005', 'pangkat_gol': 'Penata Muda (III/a)', 'jabatan': 'Statistisi Ahli Pertama'},
]

def test_gen(peg_count, blank_paras, top_margin_in=0.65, bottom_margin_in=0.55):
    doc = docx.Document(tpl_path)
    sec = doc.sections[0]
    sec.top_margin = Inches(top_margin_in)
    sec.bottom_margin = Inches(bottom_margin_in)
    
    # Update Nomor
    p_no = doc.paragraphs[2]
    p_no.text = 'Nomor: 887/SJ.9/KP.01.06/10/2026'
    p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_no.paragraph_format.space_before = Pt(0)
    p_no.paragraph_format.space_after = Pt(6)
    
    # Table
    t = doc.tables[0]
    dokumen_service.update_single_run_paragraph(
        t.rows[0].cells[2].paragraphs[0],
        'bahwa dalam rangka mendukung pelaksanaan tugas dan fungsi Pusat Kebijakan Strategis serta menindaklanjuti undangan yang diterima, perlu menugaskan pegawai untuk mengikuti kegiatan dimaksud.'
    )
    dokumen_service.update_single_run_paragraph(
        t.rows[1].cells[2].paragraphs[0],
        'Undangan Kementerian Komunikasi dan Digital RI Nomor B-1234/DJAI/HM.01.02/09/2026 tanggal 25 September 2026.'
    )
    
    plist = pegs[:peg_count]
    p1 = plist[0]
    dokumen_service.update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], 'Nama')
    dokumen_service.update_single_run_paragraph(t.rows[3].cells[3].paragraphs[0], f': {p1["nama"]}')
    dokumen_service.update_single_run_paragraph(t.rows[4].cells[2].paragraphs[0], 'NIP')
    dokumen_service.update_single_run_paragraph(t.rows[4].cells[3].paragraphs[0], f': {p1["nip"]}')
    dokumen_service.update_single_run_paragraph(t.rows[5].cells[2].paragraphs[0], 'Pangkat, Gol/Ruang')
    dokumen_service.update_single_run_paragraph(t.rows[5].cells[3].paragraphs[0], f': {p1["pangkat_gol"]}')
    dokumen_service.update_single_run_paragraph(t.rows[6].cells[2].paragraphs[0], 'Jabatan')
    dokumen_service.update_single_run_paragraph(t.rows[6].cells[3].paragraphs[0], f': {p1["jabatan"]}')
    
    curr_anchor = t.rows[6]
    for p_idx, p_data in enumerate(plist[1:], start=2):
        for field_idx in range(4):
            cloned_tr = copy.deepcopy(t.rows[3 + field_idx]._tr)
            curr_anchor._tr.addnext(cloned_tr)
            curr_anchor = docx.table._Row(cloned_tr, t)
            c_cells = curr_anchor.cells
            dokumen_service.update_single_run_paragraph(c_cells[0].paragraphs[0], '')
            dokumen_service.update_single_run_paragraph(c_cells[1].paragraphs[0], '')
            if field_idx == 0:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Nama')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["nama"]}')
            elif field_idx == 1:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'NIP')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["nip"]}')
            elif field_idx == 2:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Pangkat, Gol/Ruang')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["pangkat_gol"]}')
            elif field_idx == 3:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Jabatan')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["jabatan"]}')
                
    # Style rows
    for r_idx in range(3, len(t.rows) - 1):
        for cell in t.rows[r_idx].cells:
            dokumen_service.apply_font_to_cell(cell, sz_val=20)
            dokumen_service.set_cell_margins(cell, top=20, bottom=20, left=60, right=60)
            for p in cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                sp = pPr.find(qn('w:spacing'))
                if sp is None:
                    sp = OxmlElement('w:spacing')
                    pPr.append(sp)
                sp.set(qn('w:line'), '230')
                sp.set(qn('w:lineRule'), 'auto')
                sp.set(qn('w:before'), '0')
                sp.set(qn('w:after'), '0')
                
    # Row Untuk
    p_untuk = t.rows[-1].cells[2].paragraphs[0]
    dokumen_service.update_single_run_paragraph(
        p_untuk,
        'Menghadiri The 13th International Conference on ICT for Smart Society (ICISS) 2026 yang dilaksanakan pada 5 - 7 Oktober 2026 di Aula Timur, Institut Teknologi Bandung (Hybrid Conference).'
    )
    
    # Biaya & Demikian
    p_biaya = doc.paragraphs[3]
    p_biaya.paragraph_format.space_before = Pt(3)
    p_biaya.paragraph_format.space_after = Pt(0)
    p_demikian = doc.paragraphs[4]
    p_demikian.paragraph_format.space_before = Pt(3)
    p_demikian.paragraph_format.space_after = Pt(0)
    
    # Remove P5 (empty line before Jakarta)
    p5 = doc.paragraphs[5]
    p5._element.getparent().remove(p5._element)
    
    # Now Jakarta is doc.paragraphs[5], Jabatan is doc.paragraphs[6]
    p_tgl = doc.paragraphs[5]
    p_jab = doc.paragraphs[6]
    dokumen_service.format_signature_paragraph(p_tgl, 'Jakarta, 1 Oktober 2026', left_dxa=5200)
    dokumen_service.format_signature_paragraph(p_jab, 'Kepala Pusat Kebijakan Strategis,', left_dxa=5200)
    
    # Originally 6 empty paragraphs between Jabatan and Nama: P7..P12.
    # To leave blank_paras, remove 6 - blank_paras paragraphs.
    remove_count = 6 - blank_paras
    for idx in range(12, 12 - remove_count, -1):
        p = doc.paragraphs[idx]
        p._element.getparent().remove(p._element)
        
    p_nama = [p for p in doc.paragraphs if 'Oki' in p.text][0]
    dokumen_service.format_signature_paragraph(p_nama, 'Dr. Oki Suryowahono, M.Sc.', left_dxa=5200)
    
    out_docx = os.path.abspath(f'storage/output/test_exp_{peg_count}p_{blank_paras}blanks.docx')
    out_pdf = os.path.abspath(f'storage/output/test_exp_{peg_count}p_{blank_paras}blanks.pdf')
    doc.save(out_docx)
    dokumen_service.convert_docx_to_pdf(out_docx, out_pdf)
    r = pypdf.PdfReader(out_pdf)
    print(f'Pegawai={peg_count}, blanks={blank_paras} -> pages={len(r.pages)}')


long_dasar = 'Surat Undangan dari Deputi Bidang Kelembagaan dan Tata Laksana Kementerian PANRB Nomor: B/123/KT.01/2026 tanggal 28 September 2026 hal Undangan Rapat Koordinasi Penerapan Kebijakan SPBE Nasional.'
long_perihal = 'Menghadiri dan Mengikuti Rapat Koordinasi Tingkat Nasional Penerapan Keterpaduan Layanan Digital Nasional dan Akselerasi SPBE Terintegrasi Tahun 2026'

def test_long(blanks):
    doc = docx.Document(tpl_path)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.55)
    
    p_no = doc.paragraphs[2]
    p_no.text = 'Nomor: 887/SJ.9/KP.01.06/10/2026'
    p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_no.paragraph_format.space_before = Pt(0)
    p_no.paragraph_format.space_after = Pt(6)
    
    t = doc.tables[0]
    dokumen_service.update_single_run_paragraph(
        t.rows[0].cells[2].paragraphs[0],
        'bahwa dalam rangka mendukung pelaksanaan tugas dan fungsi Pusat Kebijakan Strategis serta menindaklanjuti undangan yang diterima, perlu menugaskan pegawai untuk mengikuti kegiatan dimaksud.'
    )
    dokumen_service.update_single_run_paragraph(t.rows[1].cells[2].paragraphs[0], long_dasar)
    
    plist = pegs[:3]
    p1 = plist[0]
    dokumen_service.update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], 'Nama')
    dokumen_service.update_single_run_paragraph(t.rows[3].cells[3].paragraphs[0], f': {p1["nama"]}')
    dokumen_service.update_single_run_paragraph(t.rows[4].cells[2].paragraphs[0], 'NIP')
    dokumen_service.update_single_run_paragraph(t.rows[4].cells[3].paragraphs[0], f': {p1["nip"]}')
    dokumen_service.update_single_run_paragraph(t.rows[5].cells[2].paragraphs[0], 'Pangkat, Gol/Ruang')
    dokumen_service.update_single_run_paragraph(t.rows[5].cells[3].paragraphs[0], f': {p1["pangkat_gol"]}')
    dokumen_service.update_single_run_paragraph(t.rows[6].cells[2].paragraphs[0], 'Jabatan')
    dokumen_service.update_single_run_paragraph(t.rows[6].cells[3].paragraphs[0], f': {p1["jabatan"]}')
    
    curr_anchor = t.rows[6]
    for p_idx, p_data in enumerate(plist[1:], start=2):
        for field_idx in range(4):
            cloned_tr = copy.deepcopy(t.rows[3 + field_idx]._tr)
            curr_anchor._tr.addnext(cloned_tr)
            curr_anchor = docx.table._Row(cloned_tr, t)
            c_cells = curr_anchor.cells
            dokumen_service.update_single_run_paragraph(c_cells[0].paragraphs[0], '')
            dokumen_service.update_single_run_paragraph(c_cells[1].paragraphs[0], '')
            if field_idx == 0:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Nama')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["nama"]}')
            elif field_idx == 1:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'NIP')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["nip"]}')
            elif field_idx == 2:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Pangkat, Gol/Ruang')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["pangkat_gol"]}')
            elif field_idx == 3:
                dokumen_service.update_single_run_paragraph(c_cells[2].paragraphs[0], 'Jabatan')
                dokumen_service.update_single_run_paragraph(c_cells[3].paragraphs[0], f': {p_data["jabatan"]}')
                
    for r_idx in range(3, len(t.rows) - 1):
        for cell in t.rows[r_idx].cells:
            dokumen_service.apply_font_to_cell(cell, sz_val=20)
            dokumen_service.set_cell_margins(cell, top=20, bottom=20, left=60, right=60)
            for p in cell.paragraphs:
                pPr = p._p.get_or_add_pPr()
                sp = pPr.find(qn('w:spacing'))
                if sp is None:
                    sp = OxmlElement('w:spacing')
                    pPr.append(sp)
                sp.set(qn('w:line'), '230')
                sp.set(qn('w:lineRule'), 'auto')
                sp.set(qn('w:before'), '0')
                sp.set(qn('w:after'), '0')
                
    p_untuk = t.rows[-1].cells[2].paragraphs[0]
    dokumen_service.update_single_run_paragraph(
        p_untuk,
        f'{long_perihal} yang dilaksanakan pada 5 - 7 Oktober 2026 di Hotel Bidakara, Jakarta Selatan.'
    )
    
    p_biaya = doc.paragraphs[3]
    p_biaya.paragraph_format.space_before = Pt(3)
    p_biaya.paragraph_format.space_after = Pt(0)
    p_demikian = doc.paragraphs[4]
    p_demikian.paragraph_format.space_before = Pt(3)
    p_demikian.paragraph_format.space_after = Pt(0)
    
    p5 = doc.paragraphs[5]
    p5._element.getparent().remove(p5._element)
    
    p_tgl = doc.paragraphs[5]
    p_jab = doc.paragraphs[6]
    dokumen_service.format_signature_paragraph(p_tgl, 'Jakarta, 1 Oktober 2026', left_dxa=5200)
    dokumen_service.format_signature_paragraph(p_jab, 'Kepala Pusat Kebijakan Strategis,', left_dxa=5200)
    
    remove_count = 6 - blanks
    for idx in range(12, 12 - remove_count, -1):
        p = doc.paragraphs[idx]
        p._element.getparent().remove(p._element)
        
    p_nama = [p for p in doc.paragraphs if 'Oki' in p.text][0]
    dokumen_service.format_signature_paragraph(p_nama, 'Dr. Oki Suryowahono, M.Sc.', left_dxa=5200)
    
    out_docx = os.path.abspath(f'storage/output/test_long_3p_{blanks}b.docx')
    out_pdf = os.path.abspath(f'storage/output/test_long_3p_{blanks}b.pdf')
    doc.save(out_docx)
    dokumen_service.convert_docx_to_pdf(out_docx, out_pdf)
    r = pypdf.PdfReader(out_pdf)
    print(f'Long text test: blanks={blanks} -> pages={len(r.pages)}')

if __name__ == '__main__':
    for b in [4, 5]:
        test_long(b)
