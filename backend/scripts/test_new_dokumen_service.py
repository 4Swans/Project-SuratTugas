import copy
import os
from datetime import date
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

NAMA_BULAN = [
    "",
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]


def format_tanggal_indonesia(d: date) -> str:
    return f"{d.day} {NAMA_BULAN[d.month]} {d.year}"


def format_periode_tanggal(tgl_mulai: date, tgl_selesai: date) -> str:
    if tgl_mulai == tgl_selesai:
        return format_tanggal_indonesia(tgl_mulai)
    if tgl_mulai.year == tgl_selesai.year:
        if tgl_mulai.month == tgl_selesai.month:
            return (
                f"{tgl_mulai.day} – {tgl_selesai.day} "
                f"{NAMA_BULAN[tgl_mulai.month]} {tgl_mulai.year}"
            )
        return (
            f"{tgl_mulai.day} {NAMA_BULAN[tgl_mulai.month]} – "
            f"{tgl_selesai.day} {NAMA_BULAN[tgl_selesai.month]} {tgl_mulai.year}"
        )
    return f"{format_tanggal_indonesia(tgl_mulai)} – {format_tanggal_indonesia(tgl_selesai)}"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f"<w:tblBorders {nsdecls('w')}>\n"
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def update_single_run_paragraph(p, new_text: str):
    """Update teks paragraf pada run pertama, kosongkan run berikutnya tanpa merusak pPr/rPr."""
    if not p.runs:
        r = p.add_run(new_text)
        r.font.name = "Arial"
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r.text = ""


def generate_surat_tugas_docx(
    template_path: str,
    output_path: str,
    nomor_surat: str,
    dasar_surat: str,
    perihal: str,
    tanggal_mulai: date,
    tanggal_selesai: date,
    lokasi: str,
    penandatangan_nama: str,
    penandatangan_jabatan: str,
    tanggal_surat: date | None = None,
    menimbang: str | None = None,
    pegawai_list: list[dict[str, Any]] | None = None,
) -> str:
    """
    Mengisi file template Word dan menyimpannya ke output_path dengan mempertahankan
    100% format naskah dinas resmi dari template-surat-tugas-blank.docx (tab, font, spasi, indents).
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template surat tidak ditemukan: {template_path}")

    pegawai_list = pegawai_list or []
    pegawai_count = len(pegawai_list)

    if tanggal_surat is None:
        tanggal_surat = date.today()

    tgl_surat_str = format_tanggal_indonesia(tanggal_surat)
    periode_str = format_periode_tanggal(tanggal_mulai, tanggal_selesai)
    tahun_surat = tanggal_mulai.year

    if not menimbang:
        menimbang = (
            "bahwa dalam rangka mendukung pelaksanaan tugas dan fungsi Pusat Kebijakan "
            "Strategis serta menindaklanjuti undangan yang diterima, perlu menugaskan "
            "pegawai untuk mengikuti kegiatan dimaksud."
        )

    doc = docx.Document(template_path)

    # Pastikan default style adalah Arial 10pt hitam jika ada run baru
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # 1. Update Nomor Surat di P2
    # Di template: Run 0='Nomor', Run 1=': ' (bold), Run 2..6='690/SJ.9/KP.01.06/08/2026'
    p_no = doc.paragraphs[2]
    if len(p_no.runs) >= 3:
        p_no.runs[0].text = "Nomor"
        p_no.runs[1].text = ": "
        p_no.runs[2].text = nomor_surat
        for r in p_no.runs[3:]:
            r.text = ""
    else:
        p_no.text = f"Nomor: {nomor_surat}"
        p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p_no.runs:
            p_no.runs[0].font.name = "Arial"

    # 2. Update Table 0 (Menimbang, Dasar, Memberi Tugas, Kepada, Untuk)
    if doc.tables:
        t = doc.tables[0]

        # Row 0: Menimbang
        p_menimbang = t.rows[0].cells[2].paragraphs[0]
        update_single_run_paragraph(p_menimbang, menimbang)

        # Row 1: Dasar
        p_dasar = t.rows[1].cells[2].paragraphs[0]
        update_single_run_paragraph(p_dasar, dasar_surat)

        # Row 3 to 6: Kepada
        if pegawai_count == 1:
            p1 = pegawai_list[0]
            # Nama
            p_nama = t.rows[3].cells[3].paragraphs[0]
            if len(p_nama.runs) >= 2:
                p_nama.runs[0].text = ": "
                p_nama.runs[1].text = p1["nama"]
                for r in p_nama.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nama, f": {p1['nama']}")

            # NIP
            p_nip = t.rows[4].cells[3].paragraphs[0]
            if len(p_nip.runs) >= 2:
                p_nip.runs[0].text = ": "
                p_nip.runs[1].text = str(p1["nip"])
                for r in p_nip.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nip, f": {p1['nip']}")

            # Pangkat, Gol/Ruang
            p_pangkat = t.rows[5].cells[3].paragraphs[0]
            p1_pangkat = p1.get("pangkat_gol") or p1.get("pangkat") or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            if len(p_pangkat.runs) >= 2:
                p_pangkat.runs[0].text = ": "
                p_pangkat.runs[1].text = str(p1_pangkat)
                for r in p_pangkat.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_pangkat, f": {p1_pangkat}")

            # Jabatan (preserves hanging indent for multiline jabatan)
            p_jabatan = t.rows[6].cells[3].paragraphs[0]
            if len(p_jabatan.runs) >= 2:
                p_jabatan.runs[0].text = ": "
                p_jabatan.runs[1].text = str(p1["jabatan"])
                for r in p_jabatan.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_jabatan, f": {p1['jabatan']}")

        elif pegawai_count in (2, 3):
            # Pegawai 1 di baris 3-6 dengan nomor urut '1. Nama'
            p1 = pegawai_list[0]
            update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], "1. Nama")
            p_nama = t.rows[3].cells[3].paragraphs[0]
            if len(p_nama.runs) >= 2:
                p_nama.runs[0].text = ": "
                p_nama.runs[1].text = p1["nama"]
                for r in p_nama.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nama, f": {p1['nama']}")

            update_single_run_paragraph(t.rows[4].cells[2].paragraphs[0], "   NIP")
            p_nip = t.rows[4].cells[3].paragraphs[0]
            if len(p_nip.runs) >= 2:
                p_nip.runs[0].text = ": "
                p_nip.runs[1].text = str(p1["nip"])
                for r in p_nip.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_nip, f": {p1['nip']}")

            update_single_run_paragraph(t.rows[5].cells[2].paragraphs[0], "   Pangkat, Gol/Ruang")
            p_pangkat = t.rows[5].cells[3].paragraphs[0]
            p1_pangkat = p1.get("pangkat_gol") or p1.get("pangkat") or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            if len(p_pangkat.runs) >= 2:
                p_pangkat.runs[0].text = ": "
                p_pangkat.runs[1].text = str(p1_pangkat)
                for r in p_pangkat.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_pangkat, f": {p1_pangkat}")

            update_single_run_paragraph(t.rows[6].cells[2].paragraphs[0], "   Jabatan")
            p_jabatan = t.rows[6].cells[3].paragraphs[0]
            if len(p_jabatan.runs) >= 2:
                p_jabatan.runs[0].text = ": "
                p_jabatan.runs[1].text = str(p1["jabatan"])
                for r in p_jabatan.runs[2:]:
                    r.text = ""
            else:
                update_single_run_paragraph(p_jabatan, f": {p1['jabatan']}")

            # Kloning baris 3-6 untuk pegawai 2 (dan 3 jika ada)
            curr_anchor = t.rows[6]
            for p_idx, p_data in enumerate(pegawai_list[1:], start=2):
                for field_idx in range(4):
                    cloned_tr = copy.deepcopy(t.rows[3 + field_idx]._tr)
                    curr_anchor._tr.addnext(cloned_tr)
                    curr_anchor = docx.table._Row(cloned_tr, t)
                    c_cells = curr_anchor.cells

                    # Kolom 0 & 1 dikosongkan (Kepada dan : hanya di pegawai pertama)
                    update_single_run_paragraph(c_cells[0].paragraphs[0], "")
                    update_single_run_paragraph(c_cells[1].paragraphs[0], "")

                    if field_idx == 0:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], f"{p_idx}. Nama")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = p_data["nama"]
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(c_cells[3].paragraphs[0], f": {p_data['nama']}")
                    elif field_idx == 1:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "   NIP")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_data["nip"])
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(c_cells[3].paragraphs[0], f": {p_data['nip']}")
                    elif field_idx == 2:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "   Pangkat, Gol/Ruang")
                        p_pkt = p_data.get("pangkat_gol") or p_data.get("pangkat") or "Penyedia Jasa Lainnya Perorangan (PJLP)"
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_pkt)
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(c_cells[3].paragraphs[0], f": {p_pkt}")
                    elif field_idx == 3:
                        update_single_run_paragraph(c_cells[2].paragraphs[0], "   Jabatan")
                        if len(c_cells[3].paragraphs[0].runs) >= 2:
                            c_cells[3].paragraphs[0].runs[0].text = ": "
                            c_cells[3].paragraphs[0].runs[1].text = str(p_data["jabatan"])
                            for r in c_cells[3].paragraphs[0].runs[2:]:
                                r.text = ""
                        else:
                            update_single_run_paragraph(c_cells[3].paragraphs[0], f": {p_data['jabatan']}")

        elif pegawai_count > 3:
            # Kasus > 3 pegawai: tulis '(Daftar Terlampir)'
            update_single_run_paragraph(t.rows[3].cells[2].paragraphs[0], "(Daftar Terlampir)")
            update_single_run_paragraph(t.rows[3].cells[3].paragraphs[0], "")
            # Hapus baris 6, 5, 4 agar tidak ada ruang kosong di tabel badan surat
            t._tbl.remove(t.rows[6]._tr)
            t._tbl.remove(t.rows[5]._tr)
            t._tbl.remove(t.rows[4]._tr)

        # Row Untuk (baris terakhir tabel 0)
        p_untuk = t.rows[-1].cells[2].paragraphs[0]
        perihal_clean = perihal.strip()
        if perihal_clean.lower().startswith("mengikuti") or perihal_clean.lower().startswith("menghadiri"):
            untuk_text = f"{perihal_clean} yang dilaksanakan pada {periode_str} di {lokasi}."
        else:
            untuk_text = f"Mengikuti kegiatan {perihal_clean} yang dilaksanakan pada {periode_str} di {lokasi}."
        update_single_run_paragraph(p_untuk, untuk_text)

    # 3. Update Paragraf Biaya, Penutup, Tanggal, dan Tanda Tangan
    # P3: Biaya
    p_biaya = doc.paragraphs[3]
    if "Tahun Anggaran" in p_biaya.text:
        p_biaya.runs[0].text = p_biaya.runs[0].text.replace("2026", str(tahun_surat))

    # P6: Tanggal Surat
    p_tgl = doc.paragraphs[6]
    if len(p_tgl.runs) >= 2:
        p_tgl.runs[0].text = "    Jakarta,"
        p_tgl.runs[1].text = f" {tgl_surat_str}"
    else:
        p_tgl.text = f"    Jakarta, {tgl_surat_str}"

    # P7: Jabatan Penandatangan
    p_jabatan_sign = doc.paragraphs[7]
    if len(p_jabatan_sign.runs) >= 2:
        p_jabatan_sign.runs[0].text = f"  {penandatangan_jabatan}"
        p_jabatan_sign.runs[1].text = ","
    else:
        p_jabatan_sign.text = f"  {penandatangan_jabatan},"

    # P14: Nama Penandatangan
    p_nama_sign = doc.paragraphs[14]
    if len(p_nama_sign.runs) >= 2:
        p_nama_sign.runs[0].text = f"         {penandatangan_nama}"
        for r in p_nama_sign.runs[1:]:
            r.text = ""
    else:
        p_nama_sign.text = f"         {penandatangan_nama}"

    # 4. Halaman Lampiran (hanya jika > 3 pegawai)
    if pegawai_count > 3:
        doc.add_page_break()

        # Header lampiran
        p_l1 = doc.add_paragraph()
        p_l1.paragraph_format.space_before = Pt(0)
        p_l1.paragraph_format.space_after = Pt(2)
        p_l1.paragraph_format.line_spacing = 1.15
        r1 = p_l1.add_run("Lampiran I Surat Tugas")
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.bold = True

        p_l2 = doc.add_paragraph()
        p_l2.paragraph_format.space_before = Pt(0)
        p_l2.paragraph_format.space_after = Pt(2)
        p_l2.paragraph_format.line_spacing = 1.15
        r2 = p_l2.add_run(f"Nomor   :  {nomor_surat}")
        r2.font.name = "Arial"
        r2.font.size = Pt(10)

        p_l3 = doc.add_paragraph()
        p_l3.paragraph_format.space_before = Pt(0)
        p_l3.paragraph_format.space_after = Pt(10)
        p_l3.paragraph_format.line_spacing = 1.15
        r3 = p_l3.add_run(f"Tanggal : {tgl_surat_str}")
        r3.font.name = "Arial"
        r3.font.size = Pt(10)

        # Tabel 5 kolom berstandar resmi
        col_widths = [Inches(0.45), Inches(1.85), Inches(1.50), Inches(1.15), Inches(1.35)]
        table_l = doc.add_table(rows=1, cols=5)
        table_l.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table_l)

        headers = ["No", "Nama", "NIP", "Pangkat/Gol\nRuang", "Jabatan"]
        hdr_cells = table_l.rows[0].cells
        for c_idx, title in enumerate(headers):
            hdr_cells[c_idx].width = col_widths[c_idx]
            set_cell_margins(hdr_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = hdr_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(title)
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.bold = True

        for idx, p_data in enumerate(pegawai_list, start=1):
            row_cells = table_l.add_row().cells
            pkt = p_data.get("pangkat_gol") or p_data.get("pangkat") or "Penyedia Jasa Lainnya Perorangan (PJLP)"
            row_vals = [
                str(idx),
                p_data["nama"],
                str(p_data["nip"]),
                pkt,
                p_data["jabatan"],
            ]
            for c_idx, val in enumerate(row_vals):
                row_cells[c_idx].width = col_widths[c_idx]
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(10)

    # Pastikan direktori output ada
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
