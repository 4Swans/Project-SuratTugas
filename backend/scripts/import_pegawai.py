"""
Script untuk mengimpor Data_Pegawai_Pusaka.xlsx ke database MySQL (tabel pegawai).
Melakukan pembersihan:
- Trim spasi berlebih di setiap kolom teks.
- Validasi NIP wajib 18 digit dan unik.
- Handling PJLP (golongan kosong -> NULL).
- Menampilkan peringatan / laporan nilai mencurigakan untuk verifikasi tim Pengolah Data.
"""

import argparse
import os
import sys
from typing import Any

import openpyxl

# Pastikan backend ada di sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.core.database import SessionLocal
from app.models.pegawai import Pegawai


def clean_str(val: Any) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def clean_nip(val: Any) -> str:
    if val is None:
        raise ValueError("NIP tidak boleh kosong")
    s = str(val).strip().replace(" ", "").replace("-", "")
    if not s.isdigit() or len(s) != 18:
        raise ValueError(f"NIP '{val}' tidak valid! Harus berupa 18 digit angka.")
    return s


def import_pegawai(file_path: str):
    if not os.path.exists(file_path):
        print(f"Error: File '{file_path}' tidak ditemukan.")
        sys.exit(1)

    print(f"Memuat file: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    headers = [clean_str(cell.value) for cell in next(ws.iter_rows())]
    print(f"Header kolom terdeteksi: {headers}")

    expected_headers = [
        "No",
        "Nama",
        "NIP",
        "Golongan",
        "Pangkat",
        "Jabatan",
        "Jurusan",
        "Pendidikan Terakhir",
        "Gelar",
        "Jenis Kelamin",
    ]
    for eh in expected_headers:
        if eh not in headers:
            print(f"Peringatan: Kolom '{eh}' tidak ditemukan di header.")

    rows = list(ws.iter_rows(values_only=True))[1:]
    print(f"Total baris data: {len(rows)}")

    db = SessionLocal()
    imported_count = 0
    updated_count = 0
    warnings: list[str] = []
    seen_nips = set()

    # Daftar keyword untuk deteksi nilai mencurigakan
    suspicious_typos = [
        "Teknik Elekto",
        "Master of Publik Administration",
        "Padjajaran",
        "(Beragama)",
    ]

    try:
        for idx, row in enumerate(rows, start=1):
            if not any(row):
                continue

            raw_nama = row[1]
            raw_nip = row[2]
            raw_golongan = row[3]
            raw_pangkat = row[4]
            raw_jabatan = row[5]
            raw_jurusan = row[6]
            raw_pendidikan = row[7]
            raw_gelar = row[8]
            raw_jk = row[9]

            nama = clean_str(raw_nama)
            if not nama:
                warnings.append(f"Baris {idx}: Nama kosong, dilewati.")
                continue

            try:
                nip = clean_nip(raw_nip)
            except ValueError as e:
                warnings.append(f"Baris {idx} ({nama}): {e}")
                continue

            if nip in seen_nips:
                raise ValueError(f"Duplikat NIP '{nip}' ditemukan pada baris {idx} ({nama})!")
            seen_nips.add(nip)

            golongan = clean_str(raw_golongan)
            pangkat = clean_str(raw_pangkat)
            jabatan = clean_str(raw_jabatan)
            jurusan = clean_str(raw_jurusan)
            pendidikan = clean_str(raw_pendidikan)
            gelar = clean_str(raw_gelar)
            jenis_kelamin = clean_str(raw_jk)

            # Deteksi nilai mencurigakan pada jurusan
            if jurusan:
                for typo in suspicious_typos:
                    if typo.lower() in jurusan.lower():
                        warnings.append(
                            f"Baris {idx} ({nama}): Indikasi typo/keterangan aneh "
                            f"di jurusan -> '{jurusan}'"
                        )

            # Deteksi anomali jenis kelamin (perlu verifikasi SK)
            if nama and jenis_kelamin:
                if (
                    "Achmad" in nama
                    and jenis_kelamin.lower() == "perempuan"
                    or "Fadhlan" in nama
                    and jenis_kelamin.lower() == "perempuan"
                ):
                    warnings.append(
                        f"Baris {idx} ({nama}): Jenis kelamin tercatat "
                        f"'{jenis_kelamin}' (perlu verifikasi SK)"
                    )

            # Cek apakah pegawai sudah ada di database
            existing = db.query(Pegawai).filter(Pegawai.nip == nip).first()
            if existing:
                existing.nama = nama
                existing.golongan = golongan
                existing.pangkat = pangkat
                existing.jabatan = jabatan
                existing.jurusan = jurusan
                existing.pendidikan_terakhir = pendidikan
                existing.gelar = gelar
                existing.jenis_kelamin = jenis_kelamin
                updated_count += 1
            else:
                new_pegawai = Pegawai(
                    nama=nama,
                    nip=nip,
                    golongan=golongan,
                    pangkat=pangkat,
                    jabatan=jabatan,
                    jurusan=jurusan,
                    pendidikan_terakhir=pendidikan,
                    gelar=gelar,
                    jenis_kelamin=jenis_kelamin,
                    status="Aktif",
                )
                db.add(new_pegawai)
                imported_count += 1

        db.commit()
        print("\n=== HASIL IMPORT PEGAWAI ===")
        print(f"Berhasil diimpor (baru): {imported_count}")
        print(f"Berhasil diperbarui (update): {updated_count}")
        print(f"Total tersimpan di DB: {imported_count + updated_count}")

        if warnings:
            print("\n=== LAPORAN PERINGATAN (UNTUK VERIFIKASI PENGOLAH DATA) ===")
            for w in warnings:
                print(f" - {w}")

    except Exception as e:
        db.rollback()
        print(f"Terjadi kesalahan saat import: {e}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import data pegawai dari Excel ke MySQL")
    parser.add_argument("file", help="Path ke file Data_Pegawai_Pusaka.xlsx")
    args = parser.parse_args()

    import_pegawai(args.file)
